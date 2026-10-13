import openpyxl
from openpyxl.utils import get_column_letter

from layout import (
    OUTPUT_COLUMNS,
    OUTPUT_FIRST_COLUMN,
    OUTPUTS_FIRST_ROW,
    OUTPUTS_SHEET,
)


def parse_output_number(value, coordinate: str) -> float:
    # LibreOffice writes integer-valued results as ints.
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    raise RuntimeError(
        f"{OUTPUTS_SHEET}!{coordinate} is {value!r}, expected a number; LibreOffice did not calculate the workbook"
    )


def load_output_rows(path, price_table: list[tuple[str, float]]) -> list[dict]:
    sheet = openpyxl.load_workbook(path, data_only=True)[OUTPUTS_SHEET]

    first_letter = get_column_letter(OUTPUT_FIRST_COLUMN)
    rows = []
    for index, (expected, _) in enumerate(price_table):
        row = OUTPUTS_FIRST_ROW + index
        found = sheet.cell(row=row, column=OUTPUT_FIRST_COLUMN).value
        if found != expected:
            raise RuntimeError(
                f"{OUTPUTS_SHEET}!{first_letter}{row} is {found!r}, expected {expected!r}; LibreOffice did not recalculate or the Outputs sheet does not line up with the price table"
            )
        item = {OUTPUT_COLUMNS[0]: found}
        for offset, key in enumerate(
            OUTPUT_COLUMNS[1:], start=OUTPUT_FIRST_COLUMN + 1
        ):
            cell = sheet.cell(row=row, column=offset)
            item[key] = parse_output_number(cell.value, cell.coordinate)
        rows.append(item)

    extra = sheet.cell(
        row=OUTPUTS_FIRST_ROW + len(price_table), column=OUTPUT_FIRST_COLUMN
    ).value
    if extra not in (None, ""):
        raise RuntimeError(
            f"{OUTPUTS_SHEET} sheet has more rows than the {len(price_table)} intervals in the price table"
        )
    return rows
